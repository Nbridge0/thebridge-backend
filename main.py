from fastapi.middleware.cors import CORSMiddleware
from fastapi import FastAPI, HTTPException, Response
from pydantic import BaseModel, EmailStr
from typing import Optional, List
from chat import (
    get_answer,
    send_help_request,
    ask_ai_only,
    save_message,
    track_click,
    add_contextual_helpful_ending
)
from supabase import create_client
import os
from dotenv import load_dotenv, find_dotenv
import secrets
import requests
from datetime import datetime, timedelta, timezone
from fastapi import UploadFile, File, Form
from openai import OpenAI

import io
import base64
import mimetypes

from pypdf import PdfReader
from docx import Document
# -------------------------
# ENV
# -------------------------
load_dotenv(find_dotenv("env.txt"))

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

supabase = create_client(SUPABASE_URL, SUPABASE_ANON_KEY)
supabase_admin = create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
openai_client = OpenAI(api_key=OPENAI_API_KEY)
# =========================================================
# CHAT ATTACHMENT HELPERS
# =========================================================

MAX_CHAT_ATTACHMENT_BYTES = 15 * 1024 * 1024

SUPPORTED_CHAT_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
}


def get_file_extension(filename: str) -> str:
    filename = str(filename or "").strip().lower()

    if "." not in filename:
        return ""

    return "." + filename.rsplit(".", 1)[-1]


def extract_chat_document_text(
    filename: str,
    content_type: str,
    file_bytes: bytes
) -> str:

    extension = get_file_extension(filename)

    # -----------------------------------------------------
    # PDF
    # -----------------------------------------------------
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))

        pages = []

        for page in reader.pages:
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""

            text = text.strip()

            if text:
                pages.append(text)

        return "\n\n".join(pages).strip()


    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------
    if extension == ".docx":
        document = Document(io.BytesIO(file_bytes))

        parts = []

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                values = [
                    (cell.text or "").strip()
                    for cell in row.cells
                ]

                row_text = " | ".join(
                    value
                    for value in values
                    if value
                )

                if row_text:
                    parts.append(row_text)

        return "\n".join(parts).strip()


    # -----------------------------------------------------
    # PLAIN TEXT / MARKDOWN / CSV
    # -----------------------------------------------------
    if extension in {".txt", ".md", ".csv"}:
        return file_bytes.decode(
            "utf-8",
            errors="ignore"
        ).strip()


    raise HTTPException(
        status_code=400,
        detail=f"Unsupported document type: {extension or content_type}"
    )


FROM_EMAIL = os.getenv("FROM_EMAIL")


def get_user_name_by_email(email: str) -> str:
    try:
        resp = supabase_admin.table("user_profiles") \
            .select("name") \
            .eq("email", email.lower().strip()) \
            .single() \
            .execute()

        if resp.data and resp.data.get("name"):
            return resp.data["name"]
    except Exception:
        pass

    return email.split("@")[0].capitalize()


def get_user_by_email(email: str):
    email = email.lower().strip()

    users = supabase_admin.auth.admin.list_users()

    for u in users:
        if u.email and u.email.lower().strip() == email:
            return u

    return None


def profile_exists(email: str) -> bool:
    email = email.lower().strip()

    try:
        resp = supabase_admin.table("user_profiles") \
            .select("id") \
            .eq("email", email) \
            .limit(1) \
            .execute()

        return bool(resp.data)
    except Exception as e:
        print("PROFILE CHECK ERROR:", e)
        return False

def track_user_signin(email: str):
    email = email.lower().strip()

    try:
        supabase_admin.table("user_signins").insert({
            "user_email": email
        }).execute()
    except Exception as e:
        print("SIGNIN TRACK ERROR:", e)


VERIFICATION_EMAIL_SUBJECT = "Almost There! Verify Your TheBridge Account"

VERIFICATION_EMAIL_BODY = """Hi {name},

Thanks for joining TheBridge — we are glad you´re here!

To finish setting up your account, please use the verification code below:

{code}

This verification code is valid for 5 minutes.

Enter the code on our website to verify your account and get started. 
If you didn´t request this, you can safely ignore this email.

Welcome aboard,
TheBridge Team
"""

WELCOME_EMAIL_SUBJECT = "Welcome to The Bridge – Your AI-Powered Superyacht Knowledge Platform"

WELCOME_EMAIL_BODY = """Hi {name},

Welcome aboard!

Thank you for signing up and joining The Bridge community. We're thrilled to have you as part of the first trusted source of verified insight built exclusively for the superyachting industry.

Your account is now active and ready for use.

If we do not have the answer yet, you have options:
• Ask AI: Get an instant reply from our integrated Open AI feature.
• Ask a Specialist: Connect with one of our industry specialists via email and receive a reply in your inbox. 
• Ask an Ambassador: Get in touch with a fellow captain or crew member who will send a reply to your inbox. 

Please allow them some time to reply.
We look forward to seeing you on The Bridge!

Kind regards,
The Bridge Team
"""


PASSWORD_RESET_EMAIL_SUBJECT = "Your Password Reset Code for TheBridge"

PASSWORD_RESET_EMAIL_BODY = """Hi {name},

We received a request to reset your TheBridge password.

Use the verification code below to continue:

{code}

This verification code is valid for 5 minutes.

Enter the code on our website to reset your password and regain access to your account.
If you didn´t request this, you can safely ignore this email.

Kind regards,
TheBridge Team
"""

HELP_EMAIL_SUBJECT = "A User From TheBridge Needs Your Help"

HELP_EMAIL_BODY = """Dear {expert_name},

We hope you are doing well!

{name} whilst using TheBridge has asked a question and selected you as their go-to {role} for help. Your insight would be greatly valued.

Question:
{question}

When you are ready, please provide a detailed answer and remember to hit Reply All.

Thank you for sharing your knowledge and helping our community grow.

Warm regards,
TheBridge Team
"""


# -------------------------
# EMAIL
# -------------------------
def send_email(to_email, subject, body):
    url = "https://api.brevo.com/v3/smtp/email"

    if isinstance(to_email, list):
        to = [{"email": e} for e in to_email]
    else:
        to = [{"email": to_email}]

    payload = {
        "sender": {
            "name": "TheBridge",
            "email": FROM_EMAIL
        },
        "to": to,
        "subject": subject,
        "textContent": body
    }

    headers = {
        "accept": "application/json",
        "api-key": os.getenv("BREVO_API_KEY"),
        "content-type": "application/json"
    }

    response = requests.post(url, json=payload, headers=headers, timeout=10)

    if response.status_code >= 400:
        print("❌ BREVO API ERROR:", response.status_code, response.text)
        raise Exception("Email send failed")
# -------------------------
# APP
# -------------------------
app = FastAPI(title="TheBridge API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,  # ✅ correct for ngrok/browser
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/experts")
def list_experts(role: str):
    experts = supabase_admin.table("experts") \
        .select("id, name, email, description") \
        .eq("role", role) \
        .eq("is_active", True) \
        .execute()

    return experts.data


@app.post("/transcribe")
async def transcribe_audio(file: UploadFile = File(...)):
    try:
        # Read file bytes
        contents = await file.read()

        # Re-wrap properly with filename + content type
        transcription = openai_client.audio.transcriptions.create(
            model="gpt-4o-mini-transcribe",
            file=(
                file.filename,
                contents,
                file.content_type
            )
        )

        return {"text": transcription.text}

    except Exception as e:
        print("TRANSCRIPTION ERROR:", e)
        raise HTTPException(status_code=500, detail="Transcription failed")


# -------------------------
# MODELS
# -------------------------


class SignupRequest(BaseModel):
    name: str
    email: EmailStr
    password: str
    newsletter: bool = False

class LoginRequest(BaseModel):
    email: EmailStr
    password: str


class VerifyRequest(BaseModel):
    email: EmailStr
    code: str


class ResetRequest(BaseModel):
    email: EmailStr


class ResetVerifyRequest(BaseModel):
    email: EmailStr
    code: str


class ResetConfirmRequest(BaseModel):
    email: EmailStr
    new_password: str

class DeleteAccountRequest(BaseModel):
    email: EmailStr

class CreateChatRequest(BaseModel):
    user_email: EmailStr
    title: str

class HelpRequest(BaseModel):
    message: str
    user_email: EmailStr
    expert_emails: List[EmailStr]
    role: str                     # 'specialist' or 'ambassador'
    chat_id: Optional[int] = None
    user_role: str = "guest"


class RenameChatRequest(BaseModel):
    chat_id: int
    title: str

class ChatRequest(BaseModel):
    chat_id: Optional[int] = None
    message: str
    user_role: str = "guest"
    user_email: Optional[str] = None
    history: Optional[list] = []

class AnswerActionRequest(BaseModel):
    chat_id: Optional[int] = None
    user_email: Optional[str] = None
    user_role: str = "guest"
    action: str
    answer_text: Optional[str] = None
    source: Optional[str] = None


class TTSRequest(BaseModel):
    text: str

class UpdateProfileRequest(BaseModel):
    current_email: EmailStr
    name: str
    email: EmailStr

@app.post("/answer/action")
def save_answer_action(req: AnswerActionRequest):
    allowed_actions = {"share", "good", "bad", "read"}

    if req.action not in allowed_actions:
        raise HTTPException(status_code=400, detail="Invalid action")

    try:
        supabase_admin.table("answer_actions").insert({
            "chat_id": req.chat_id,
            "user_email": req.user_email,
            "user_type": "user" if req.user_role != "guest" else "guest",
            "action": req.action,
            "answer_text": req.answer_text,
            "source": req.source
        }).execute()

        return {"status": "saved"}

    except Exception as e:
        print("ANSWER ACTION SAVE ERROR:", e)
        raise HTTPException(status_code=500, detail="Failed to save answer action")

@app.post("/tts")
def text_to_speech(req: TTSRequest):
    text = (req.text or "").strip()

    if not text:
        raise HTTPException(status_code=400, detail="Text is required")

    # Keep requests safe and fast
    if len(text) > 4000:
        text = text[:4000]

    try:
        speech = openai_client.audio.speech.create(
            model="gpt-4o-mini-tts",
            voice="nova",
            input=text,
            response_format="mp3",
            instructions=(
                "Speak clearly in a warm, professional, female-sounding voice. "
                "Automatically pronounce the input in its original language. "
                "Keep the tone calm, natural, and easy to understand."
            )
        )

        return Response(
            content=speech.content,
            media_type="audio/mpeg"
        )

    except Exception as e:
        print("TTS ERROR:", e)
        raise HTTPException(status_code=500, detail="Text to speech failed")



@app.put("/chats/{chat_id}/rename")
def rename_chat(chat_id: int, payload: dict):

    title = payload.get("title")
    user_email = payload.get("user_email")

    if not title:
        raise HTTPException(status_code=400, detail="Title required")

    resp = supabase_admin.table("user_chats") \
        .update({"title": title}) \
        .eq("id", chat_id) \
        .eq("user_email", user_email) \
        .execute()

    if not resp.data:
        raise HTTPException(status_code=404, detail="Chat not found")

    return {"status": "renamed"}

@app.post("/help/send")
def help_send(req: HelpRequest):

    # ✅ TRACK BUTTON CLICK
    track_click(
        chat_id=req.chat_id,
        button="ask_specialist" if req.role == "specialist" else "ask_ambassador",
        question=req.message,
        user_email=req.user_email,
        user_role=req.user_role
    )


    # 🔒 Re-fetch valid experts by role
    valid_experts = supabase_admin.table("experts") \
        .select("email") \
        .eq("role", req.role) \
        .eq("is_active", True) \
        .in_("email", req.expert_emails) \
        .execute()

    if not valid_experts.data:
        raise HTTPException(400, "Invalid expert selection")

    for expert in valid_experts.data:
        send_help_request(
            req.role,
            req.message,
            req.user_email,
            expert["email"]
        )

    return {"status": "emails_sent"}


# -------------------------
# HEALTH
# -------------------------
@app.get("/health")
def health():
    return {"status": "ok"}

@app.post("/chat/attachment")
async def chat_attachment(
    file: UploadFile = File(...),
    message: str = Form(""),
    chat_id: Optional[int] = Form(None),
    user_role: str = Form("guest"),
    user_email: Optional[str] = Form(None)
):
    """
    Analyse a picture or document attached directly to a chat.

    This route is intentionally separate from /chat/message so
    the existing normal chat routing remains unchanged.
    """

    filename = (file.filename or "attachment").strip()
    content_type = (
        file.content_type
        or mimetypes.guess_type(filename)[0]
        or "application/octet-stream"
    )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    if len(file_bytes) > MAX_CHAT_ATTACHMENT_BYTES:
        raise HTTPException(
            status_code=413,
            detail="The uploaded file is too large."
        )

    user_question = (message or "").strip()

    if not user_question:
        user_question = (
            "Analyse this attachment and explain the relevant "
            "information clearly."
        )
        
    track_click(
        chat_id=chat_id,
        button="question",
        question=user_question,
        user_email=user_email,
        user_role=user_role
    )


    # =====================================================
    # CREATE CHAT IF REQUIRED
    # =====================================================

    effective_chat_id = chat_id

    if effective_chat_id is None and user_email:

        new_chat = (
            supabase_admin
            .table("user_chats")
            .insert({
                "user_email": user_email,
                "title": user_question[:40] or filename[:40]
            })
            .execute()
        )

        if new_chat.data:
            effective_chat_id = new_chat.data[0]["id"]


    # =====================================================
    # IMAGE
    # =====================================================

    if content_type.startswith("image/"):

        encoded = base64.b64encode(
            file_bytes
        ).decode("utf-8")

        data_url = (
            f"data:{content_type};base64,{encoded}"
        )

        response = openai_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are TheBridge AI. "
                        "Analyse the user-provided image carefully. "
                        "Answer the user's actual question about the image. "
                        "Do not invent details that are not visible or "
                        "reasonably supported by the image. "
                        "If something cannot be determined from the image, "
                        "say so clearly."
                    )
                },
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": user_question
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": data_url
                            }
                        }
                    ]
                }
            ],
            temperature=0
        )

        answer = (
            response
            .choices[0]
            .message
            .content
            .strip()
        )
        answer = add_contextual_helpful_ending(
            user_question,
            answer
        )

        source = "chat_image_attachment"


    # =====================================================
    # DOCUMENT
    # =====================================================

    else:

        extracted_text = extract_chat_document_text(
            filename=filename,
            content_type=content_type,
            file_bytes=file_bytes
        )

        if not extracted_text:
            raise HTTPException(
                status_code=400,
                detail=(
                    "No readable text could be extracted "
                    "from this document."
                )
            )

        # Prevent a huge attachment from overflowing one request.
        extracted_text = extracted_text[:120000]

        response = openai_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are TheBridge AI. "
                        "Answer the user's question using the attached "
                        "document content supplied below. "
                        "Do not invent information that is not supported "
                        "by the attachment. "
                        "If the attachment does not contain the answer, "
                        "say so clearly."
                    )
                },
                {
                    "role": "user",
                    "content": (
                        f"User question:\n"
                        f"{user_question}\n\n"
                        f"Attached filename:\n"
                        f"{filename}\n\n"
                        f"Attached document content:\n"
                        f"{extracted_text}"
                    )
                }
            ],
            temperature=0
        )

        answer = (
            response
            .choices[0]
            .message
            .content
            .strip()
        )

        answer = add_contextual_helpful_ending(
            user_question,
            answer
        )

        source = "chat_document_attachment"


    # =====================================================
    # SAVE TO CHAT HISTORY
    # =====================================================

    if effective_chat_id is not None:

        save_message(
            effective_chat_id,
            "user",
            f"{user_question}\n\n[Attachment: {filename}]",
            "user_attachment",
            user_email
        )

        save_message(
            effective_chat_id,
            "assistant",
            answer,
            source,
            user_email
        )


    return {
        "answer": answer,
        "source": source,
        "actions": [],
        "requires_auth": False,
        "chat_id": effective_chat_id,
        "filename": filename
    }
    
# -------------------------
# CHAT
# -------------------------
@app.post("/chat/message")
def chat_message(req: ChatRequest):

    # =====================================================
    # TRACK EVERY QUESTION — USERS + GUESTS
    # =====================================================
    track_click(
        chat_id=req.chat_id,
        button="question",
        question=req.message,
        user_email=req.user_email,
        user_role=req.user_role
    )

    # ✅ FIX: Ensure chat exists (important for suggested questions)
    if req.chat_id is None and req.user_email:
        new_chat = supabase_admin.table("user_chats").insert({
            "user_email": req.user_email,
            "title": "New Chat"
        }).execute()

        req.chat_id = new_chat.data[0]["id"]

    # Save user message
    if req.chat_id is not None:
        save_message(
            req.chat_id,
            "user",
            req.message,
            "user",
            req.user_email
        )

    try:
        result = get_answer(req.message, req.user_role, req.chat_id, req.history)
    except Exception as e:
        print("AI ERROR:", e)
        return {
            "answer": "⚠️ Temporary error. Please try again.",
            "source": "error",
            "actions": ["ask_ai"],
            "requires_auth": False
        }

    # ✅ MULTI-PARTNER SUPPORT
    if "answers" in result:
        if req.chat_id is not None:
            for ans in result["answers"]:
                save_message(
                    req.chat_id,
                    "assistant",
                    ans.get("answer"),
                    result.get("source"),
                    req.user_email,
                    ans.get("partner_name")
                )
        return result

    answer = result.get("answer")
    source = result.get("source")
    actions = result.get("actions", [])
    requires_auth = result.get("requires_auth", False)

    if req.chat_id is not None:
        save_message(
            req.chat_id,
            "assistant",
            answer,
            source,
            req.user_email,
            result.get("badge")
        )
    return {
    "answer": answer,
    "source": source,
    "actions": actions,
    "requires_auth": requires_auth,
    "new_title": result.get("new_title")
    }

@app.post("/chat/ask-ai")
def chat_ask_ai(req: dict):
    """
    Handles:
    - normal Ask AI flow
    - frontend-only click tracking for gated buttons
    """

    # 🔒 FRONTEND-ONLY CLICK TRACKING (guest intent)
    if req.get("__track_only") is True:
        track_click(
            chat_id=req.get("chat_id"),
            button=req.get("button_override"),
            question=req.get("message"),
            user_email=req.get("user_email"),
            user_role=req.get("user_role", "guest")
        )
        return {"status": "tracked"}

    # ✅ NORMAL Ask AI FLOW (unchanged behavior)
    track_click(
        chat_id=req.get("chat_id"),
        button="ask_ai",
        question=req.get("message"),
        user_email=req.get("user_email"),
        user_role=req.get("user_role", "guest")
    )

    answer = ask_ai_only(
         req.get("message"),
         req.get("chat_id"),
         req.get("history")
    )


    if req.get("chat_id"):
        save_message(
            req.get("chat_id"),
            "assistant",
            answer,
            "openai_only",
            req.get("user_email")
        )

    return {
        "answer": answer,
        "source": "openai_only"
    }

# -------------------------
# AUTH
# -------------------------

@app.get("/auth/profile")
def get_profile(email: EmailStr):
    email = email.lower().strip()

    try:
        profile = (
            supabase_admin
            .table("user_profiles")
            .select("id, name, email, newsletter")
            .eq("email", email)
            .single()
            .execute()
        )

        if not profile.data:
            raise HTTPException(
                status_code=404,
                detail="Profile not found"
            )

        return profile.data

    except HTTPException:
        raise

    except Exception as e:
        print("GET PROFILE ERROR:", repr(e))
        raise HTTPException(
            status_code=500,
            detail="Failed to load profile"
        )


@app.put("/auth/profile")
def update_profile(req: UpdateProfileRequest):
    current_email = req.current_email.lower().strip()
    new_email = req.email.lower().strip()
    new_name = req.name.strip()

    if not new_name:
        raise HTTPException(
            status_code=400,
            detail="Name is required"
        )

    try:
        user = get_user_by_email(current_email)

        if not user:
            raise HTTPException(
                status_code=404,
                detail="User not found"
            )

        # If email is changing, check that new email is not already used
        if new_email != current_email:
            existing_user = get_user_by_email(new_email)

            if existing_user:
                raise HTTPException(
                    status_code=409,
                    detail="This email is already in use"
                )

        # 1. Update Supabase Auth user
        supabase_admin.auth.admin.update_user_by_id(
            user.id,
            {
                "email": new_email,
                "email_confirm": True,
                "user_metadata": {
                    "name": new_name
                }
            }
        )

        # 2. Update profile
        (
            supabase_admin
            .table("user_profiles")
            .update({
                "name": new_name,
                "email": new_email
            })
            .eq("email", current_email)
            .execute()
        )

        # 3. Keep chats attached to the new email
        (
            supabase_admin
            .table("user_chats")
            .update({"user_email": new_email})
            .eq("user_email", current_email)
            .execute()
        )

        # 4. Keep messages attached to the new email
        (
            supabase_admin
            .table("chat_messages")
            .update({"user_email": new_email})
            .eq("user_email", current_email)
            .execute()
        )

        # 5. Keep click tracking attached to the new email
        (
            supabase_admin
            .table("user_clicks")
            .update({"user_email": new_email})
            .eq("user_email", current_email)
            .execute()
        )

        # 6. Update pending reset / verification rows if any exist
        (
            supabase_admin
            .table("password_resets")
            .update({"email": new_email})
            .eq("email", current_email)
            .execute()
        )

        (
            supabase_admin
            .table("email_verifications")
            .update({"email": new_email})
            .eq("email", current_email)
            .execute()
        )

        return {
            "status": "updated",
            "name": new_name,
            "email": new_email
        }

    except HTTPException:
        raise

    except Exception as e:
        print("UPDATE PROFILE ERROR:", repr(e))
        raise HTTPException(
            status_code=500,
            detail="Profile update failed"
        )

@app.post("/auth/signup")
def signup(req: SignupRequest):
    email = req.email.lower().strip()

    try:
        existing_user = get_user_by_email(email)

        if existing_user or profile_exists(email):
            raise HTTPException(
                status_code=409,
                detail="Account already exists. Please log in instead."
            )

        code = secrets.token_hex(3)
        expiry = datetime.now(timezone.utc) + timedelta(minutes=5)

        upsert_resp = (
            supabase_admin
            .table("email_verifications")
            .upsert({
                "email": email,
                "name": req.name,
                "password": req.password,
                "newsletter": req.newsletter,
                "code": code,
                "expires_at": expiry.isoformat()
            })
            .execute()
        )

        if not upsert_resp.data:
            raise HTTPException(
                status_code=500,
                detail="Failed to save verification code"
            )

        send_email(
            email,
            VERIFICATION_EMAIL_SUBJECT,
            VERIFICATION_EMAIL_BODY.format(
                name=req.name,
                code=code
            )
        )

        return {"status": "verification_sent"}

    except HTTPException:
        raise

    except Exception as e:
        print("SIGNUP ERROR:", repr(e))
        raise HTTPException(
            status_code=500,
            detail="Signup failed"
        )
@app.post("/auth/verify")
def verify(req: VerifyRequest):
    email = req.email.lower().strip()
    code = req.code.strip()

    try:
        # 1. Fetch verification record
        resp = (
            supabase_admin
            .table("email_verifications")
            .select("*")
            .eq("email", email)
            .limit(1)
            .execute()
        )

        if not resp.data:
            raise HTTPException(
                status_code=400,
                detail="No verification request found"
            )

        record = resp.data[0]

        # 2. Check expiry safely
        expires_at = datetime.fromisoformat(
            record["expires_at"].replace("Z", "+00:00")
        )

        if expires_at < datetime.now(timezone.utc):
            raise HTTPException(
                status_code=400,
                detail="Verification code expired"
            )

        # 3. Check code
        if code != str(record["code"]).strip():
            raise HTTPException(
                status_code=400,
                detail="Invalid verification code"
            )

        # 4. Check if auth user already exists
        existing_user = get_user_by_email(email)

        if existing_user:
            user_id = existing_user.id
        else:
            # 5. Create user with ADMIN client
            created_user = supabase_admin.auth.admin.create_user({
                "email": email,
                "password": record["password"],
                "email_confirm": True,
                "user_metadata": {
                    "name": record["name"]
                }
            })

            if not created_user or not created_user.user:
                raise HTTPException(
                    status_code=500,
                    detail="Failed to create user account"
                )

            user_id = created_user.user.id

        # 6. Create/update profile
        profile_resp = (
            supabase_admin
            .table("user_profiles")
            .upsert({
                "id": user_id,
                "email": email,
                "name": record["name"],
                "newsletter": record.get("newsletter", False)
            })
            .execute()
        )

        if not profile_resp.data:
            raise HTTPException(
                status_code=500,
                detail="Failed to create user profile"
            )

        # 7. Delete verification record ONLY after success
        (
            supabase_admin
            .table("email_verifications")
            .delete()
            .eq("email", email)
            .execute()
        )

        # 8. Send welcome email, but don't fail verification if email fails
        try:
            send_email(
                email,
                WELCOME_EMAIL_SUBJECT,
                WELCOME_EMAIL_BODY.format(name=record["name"])
            )
        except Exception as e:
            print("WELCOME EMAIL ERROR:", e)

        return {
            "status": "verified",
            "user_id": user_id
        }

    except HTTPException:
        raise

    except Exception as e:
        print("VERIFY ERROR:", repr(e))
        raise HTTPException(
            status_code=500,
            detail="Server error during verification"
        )

@app.post("/auth/login")
def login(req: LoginRequest):
    email = req.email.lower().strip()

    if not profile_exists(email):
        raise HTTPException(
            status_code=401,
            detail="Account not found. Please create an account again."
        )

    try:
        auth = supabase.auth.sign_in_with_password({
            "email": email,
            "password": req.password
        })

        track_user_signin(email)

        profile = (
            supabase_admin
            .table("user_profiles")
            .select("name, email")
            .eq("email", email)
            .single()
            .execute()
        )

        return {
            "status": "ok",
            "user": auth.user,
            "name": profile.data.get("name") if profile.data else None,
            "email": email
        }

    except Exception:
        raise HTTPException(status_code=401, detail="Invalid credentials")
        
@app.delete("/auth/account")
def delete_user_account(req: DeleteAccountRequest):
    email = req.email.lower().strip()

    try:
        # 1. Find Supabase Auth user first
        user = get_user_by_email(email)

        # 2. Get this user's chats
        chats = (
            supabase_admin
            .table("user_chats")
            .select("id")
            .eq("user_email", email)
            .execute()
        )

        chat_ids = [chat["id"] for chat in chats.data or []]

        # 3. Delete chat messages for those chats
        if chat_ids:
            (
                supabase_admin
                .table("chat_messages")
                .delete()
                .in_("chat_id", chat_ids)
                .execute()
            )

        # 4. Delete user's chats
        (
            supabase_admin
            .table("user_chats")
            .delete()
            .eq("user_email", email)
            .execute()
        )

        # 5. Delete click analytics for this user
        (
            supabase_admin
            .table("user_clicks")
            .delete()
            .eq("user_email", email)
            .execute()
        )

        # 6. Delete any pending password reset rows
        (
            supabase_admin
            .table("password_resets")
            .delete()
            .eq("email", email)
            .execute()
        )

        # 7. Delete any pending verification rows
        (
            supabase_admin
            .table("email_verifications")
            .delete()
            .eq("email", email)
            .execute()
        )

        # 8. Delete user profile
        (
            supabase_admin
            .table("user_profiles")
            .delete()
            .eq("email", email)
            .execute()
        )

        # 9. Delete Supabase Auth user
        if user:
            supabase_admin.auth.admin.delete_user(user.id)

        return {
            "status": "deleted",
            "email": email
        }

    except Exception as e:
        print("DELETE ACCOUNT ERROR:", repr(e))
        raise HTTPException(
            status_code=500,
            detail="Account deletion failed"
        )
# -------------------------
# PASSWORD RESET (BASIC)
# -------------------------
@app.post("/auth/password-reset/request")
def reset_request(req: ResetRequest):
    email = req.email.lower().strip()

    user = get_user_by_email(email)
    if not user:
        raise HTTPException(
            status_code=404,
            detail="Email not registered with us. Please create an account."
        )

    code = secrets.token_hex(3)
    expiry = datetime.now(timezone.utc) + timedelta(minutes=5)

    supabase_admin.table("password_resets").upsert({
        "email": email,
        "code": code,
        "expires_at": expiry.isoformat()
    }).execute()

    name = get_user_name_by_email(email)

    send_email(
        email,
        PASSWORD_RESET_EMAIL_SUBJECT,
        PASSWORD_RESET_EMAIL_BODY.format(name=name, code=code)
        
    )

    return {"status": "code_sent"}


@app.post("/auth/password-reset/verify")
def reset_verify(req: ResetVerifyRequest):
    resp = supabase_admin.table("password_resets") \
        .select("*") \
        .eq("email", req.email.lower()) \
        .single() \
        .execute()

    if not resp.data:
        raise HTTPException(400, "No reset request found")

    record = resp.data
    expires_at = datetime.fromisoformat(record["expires_at"])

    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(400, "Code expired")

    if req.code != record["code"]:
        raise HTTPException(400, "Invalid code")

    return {"status": "verified"}



@app.post("/auth/password-reset/confirm")
def reset_confirm(req: ResetConfirmRequest):
    email = req.email.lower().strip()

    # 1. Ensure reset was requested
    reset = supabase_admin.table("password_resets") \
        .select("*") \
        .eq("email", email) \
        .single() \
        .execute()

    if not reset.data:
        raise HTTPException(400, "Reset not verified")

    expires_at = datetime.fromisoformat(reset.data["expires_at"])
    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(400, "Reset expired")

    # 2. Update password
    user = get_user_by_email(email)
    if not user:
         raise HTTPException(404, "User not found")


    supabase_admin.auth.admin.update_user_by_id(
        user.id,
        {
         "password": req.new_password,
         "email_confirmed_at": datetime.now(timezone.utc).isoformat()
        }
    )

    # 3. DELETE reset record (ONLY HERE)
    supabase_admin.table("password_resets") \
        .delete() \
        .eq("email", email) \
        .execute()

    return {"status": "password_updated"}



# -------------------------
# CHATS
# -------------------------
@app.get("/chats")
def list_chats(user_email: EmailStr):
    resp = supabase_admin.table("user_chats") \
        .select("*") \
        .eq("user_email", user_email) \
        .execute()
    return resp.data 


@app.post("/chats")
def create_chat(payload: dict):
    user_email = payload.get("user_email")
    title = payload.get("title") or "New Chat"

    # Only enforce email if user is logged in
    if not user_email:
        # Guests don't get persistent chats
        return {"chat_id": None}

    result = (
        supabase_admin
        .table("user_chats")
        .insert({
            "user_email": user_email,
            "title": title
        })
        .execute()
    )

    if not result.data:
        raise HTTPException(
            status_code=500,
            detail="Failed to create chat"
        )

    return {
        "chat_id": result.data[0]["id"]
    }


@app.get("/chats/{chat_id}/messages")
def get_chat_messages(chat_id: int):
    resp = supabase_admin.table("chat_messages") \
        .select("*") \
        .eq("chat_id", chat_id) \
        .order("id") \
        .execute()
    return resp.data


@app.delete("/chats/{chat_id}")
def delete_chat(chat_id: int, user_email: str):
    supabase_admin.table("chat_messages") \
        .delete() \
        .eq("chat_id", chat_id) \
        .execute()

    supabase_admin.table("user_chats") \
        .delete() \
        .eq("id", chat_id) \
        .eq("user_email", user_email) \
        .execute()

    return {"status": "deleted"}


# -------------------------
# SUGGESTED QUESTIONS
# -------------------------
@app.get("/suggested-questions")
def get_suggested_questions():

    resp = supabase_admin.table("suggested_questions") \
        .select("question") \
        .eq("is_active", True) \
        .order("display_order") \
        .execute()

    if not resp.data:
        return []

    return [row["question"] for row in resp.data]